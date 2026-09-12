"""Teacher-facing download projection; never modifies a saved schedule."""
from io import BytesIO
from datetime import date, datetime
import re

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment


def _label(value, names):
    raw = str(value or "").strip()
    if not raw:
        return ""
    resolved = str(names.get(raw, "") or "").strip()
    if resolved:
        return resolved
    if re.search(r"[0-9a-f]{8}-[0-9a-f-]{20,}|^(subject|class|component|lesson)-|^SUB-", raw, re.I):
        return "Chưa xác định"
    return raw


def localize_download(content, *, class_names, subject_names, component_names,
                      teacher_profile=None, sessions=None):
    """Return new XLSX bytes. Preserve dates, periods, titles and equipment."""
    workbook = load_workbook(BytesIO(content))
    try:
        sheet = workbook["Lich_bao_giang"]
        if sheet["A5"].value != "Ngày dạy" or sheet["H5"].value != "Mã bài":
            raise ValueError("Mẫu Excel đã thay đổi; chưa thể Việt hóa an toàn.")
        profile = teacher_profile or {}
        old_heading = str(sheet["A2"].value or "")
        year = old_heading.split("Năm học:", 1)
        if len(year) != 2:
            raise ValueError("Không tìm thấy năm học trong file xuất.")
        parts = []
        if profile.get("show_teacher_name", True):
            parts.append("Giáo viên: " + str(profile.get("full_name") or "Chưa xác định"))
        if profile.get("show_school_name", True) and profile.get("school_name"):
            parts.append("Trường: " + str(profile["school_name"]))
        parts.append("Năm học: " + year[1].strip())
        sheet["A2"] = "  |  ".join(parts)
        sheet["H5"] = "Buổi"
        for table in sheet.tables.values():
            for column in table.tableColumns:
                if column.name == "Mã bài":
                    column.name = "Buổi"
        for cells in sheet.iter_rows(min_row=6):
            if not isinstance(cells[0].value, (date, datetime)):
                continue
            day = cells[0].value
            day_key = day.date() if isinstance(day, datetime) else day
            key = (day_key, str(cells[3].value), cells[2].value,
                   str(cells[4].value), cells[6].value)
            session = (sessions or {}).get(key, "")
            for index, names in ((3, class_names), (4, subject_names), (5, component_names)):
                cells[index].value = _label(cells[index].value, names)
            cells[7].value = {"MORNING": "Sáng", "AFTERNOON": "Chiều"}.get(session, "Chưa xác định")
            cells[0].number_format = "dd/mm/yyyy"
        sheet.title = "Lịch báo giảng"
        for row in sheet:
            for cell in row:
                if cell.value is not None:
                    cell.font = Font(name="Times New Roman", size=12,
                                     bold=cell.row in (1, 2, 5),
                                     color="FFFFFF" if cell.row in (1, 5) else "000000")
                    if isinstance(cell.value, str):
                        cell.data_type = "s"
        sheet["A1"].fill = PatternFill("solid", fgColor="16324F")
        for cell in sheet[5]:
            cell.fill = PatternFill("solid", fgColor="2364AA")
        sheet["A2"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.row_dimensions[2].height = 36
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()
    finally:
        workbook.close()

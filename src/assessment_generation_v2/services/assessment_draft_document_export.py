"""Read only export of one validated exam draft and its approved blueprint."""

from __future__ import annotations

from collections import Counter
from decimal import Decimal
from io import BytesIO
from uuid import UUID

from docx import Document
from docx.shared import RGBColor
from docx.shared import Cm, Pt


def _rows(query):
    data = query.execute().data
    if not isinstance(data, list):
        raise ValueError("Dữ liệu xuất đề không hợp lệ.")
    return data


def _one(query, label):
    rows = _rows(query)
    if len(rows) != 1:
        raise ValueError(f"Không tìm thấy đúng một {label} được phép xem.")
    return rows[0]


def _table(doc, headings, records):
    table = doc.add_table(rows=1, cols=len(headings))
    table.style = "Table Grid"
    for cell, heading in zip(table.rows[0].cells, headings):
        cell.text = heading
    for record in records:
        for cell, value in zip(table.add_row().cells, record):
            cell.text = str(value)


def _doc(title, subtitle):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.space_after = Pt(4)
    title_style = doc.styles["Title"]
    title_style.font.name = "Times New Roman"
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_style.font.size = Pt(16)
    paragraph = doc.add_paragraph(title, "Title")
    for element in paragraph._p.xpath("./w:pPr/w:pBdr"):
        element.getparent().remove(element)
    for element in title_style.element.xpath("./w:pPr/w:pBdr"):
        element.getparent().remove(element)
    doc.add_paragraph(subtitle)
    return doc


def _bytes(doc):
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def export_draft_documents(*, client, user_id: str, exam_version_id: str):
    """Return filename/bytes for exam, matrix, specification, answer key.

    All lookups use the caller's Supabase session and enforce ownership locally.
    The export never changes review status or creates a publication snapshot.
    """
    owner = str(UUID(user_id))
    version_id = str(UUID(exam_version_id))
    version = _one(
        client.table("assessment_exam_versions").select(
            "exam_version_id,exam_title,exam_code_label,total_score,duration_minutes,"
            "assembly_status,blueprint_version_id,assessment_exams!inner(owner_user_id,grade_level,exam_code)"
        ).eq("exam_version_id", version_id).eq("assessment_exams.owner_user_id", owner),
        "bản đề",
    )
    exam = version["assessment_exams"]
    if isinstance(exam, list):
        exam = exam[0] if len(exam) == 1 else None
    if not isinstance(exam, dict) or str(exam.get("owner_user_id")) != owner:
        raise PermissionError("Không có quyền xuất bản đề này.")
    if version["assembly_status"] not in ("ASSEMBLED", "PENDING_REVIEW"):
        raise ValueError("Bản đề chưa được lắp ráp để xuất bản nháp.")

    blueprint_id = version["blueprint_version_id"]
    blueprint = _one(client.table("assessment_blueprint_versions").select(
        "blueprint_version_id,review_status,blueprint_name,total_score,"
        "assessment_blueprints!inner(owner_user_id)"
    ).eq("blueprint_version_id", blueprint_id).eq("assessment_blueprints.owner_user_id", owner), "ma trận")
    if blueprint["review_status"] != "APPROVED":
        raise ValueError("Ma trận chưa được duyệt.")
    cells = _rows(client.table("assessment_blueprint_cells").select(
        "blueprint_cell_id,section_code,topic_code,cognitive_level_code,"
        "question_type_code,question_count,response_count,target_score,specification_note,sequence_number"
    ).eq("blueprint_version_id", blueprint_id).order("sequence_number"))
    assignments = _rows(client.table("assessment_exam_questions").select(
        "display_number,blueprint_cell_id,question_version_id,assigned_score,"
        "assessment_question_versions!inner(prompt_text,question_type_code,review_status,"
        "assessment_question_options(option_code,option_text,sequence_number,is_correct),"
        "assessment_question_statements(statement_code,statement_text,sequence_number,correct_value),"
        "assessment_question_answers(exact_answer_text,numeric_answer,answer_explanation),"
        "assessment_question_solutions(solution_text,sequence_number,is_primary),"
        "assessment_question_requirement_links(requirement_code,link_role))"
    ).eq("exam_version_id", version_id).order("display_number"))
    if not cells or not assignments:
        raise ValueError("Ma trận hoặc câu hỏi lắp ráp còn thiếu.")
    positions = [int(row["display_number"]) for row in assignments]
    if positions != list(range(1, len(assignments) + 1)):
        raise ValueError("Số thứ tự câu hỏi không liên tục.")
    cell_map = {row["blueprint_cell_id"]: row for row in cells}
    counts = Counter(row["blueprint_cell_id"] for row in assignments)
    if any(counts.get(cell_id) != int(cell["question_count"]) for cell_id, cell in cell_map.items()) or set(counts) != set(cell_map):
        raise ValueError("Số câu theo ô không khớp ma trận.")
    if sum((Decimal(str(row["assigned_score"])) for row in assignments), Decimal(0)) != Decimal(str(version["total_score"])):
        raise ValueError("Điểm của câu hỏi không khớp tổng điểm đề.")
    if Decimal(str(blueprint["total_score"])) != Decimal(str(version["total_score"])):
        raise ValueError("Điểm của ma trận không khớp bản đề.")
    for row in assignments:
        question = row["assessment_question_versions"]
        if question["review_status"] != "APPROVED" or not str(question["prompt_text"]).strip():
            raise ValueError("Đề chứa câu hỏi chưa được duyệt hoặc chưa có nội dung.")

    requirements = sorted({link["requirement_code"] for row in assignments for link in row["assessment_question_versions"].get("assessment_question_requirement_links", [])})
    requirement_text = {}
    if requirements:
        requirement_text = {r["requirement_code"]: r["requirement_text"] for r in _rows(client.table("assessment_learning_requirements").select("requirement_code,requirement_text").in_("requirement_code", requirements))}
    subtitle = (f"BẢN NHÁP CẦN DUYỆT · Lớp {exam['grade_level']} · "
                f"{version['duration_minutes']} phút · {version['total_score']} điểm · "
                f"Mã đề: {version['exam_code_label'] or exam['exam_code']}")
    exam_doc = _doc(version["exam_title"], subtitle)
    answer_doc = _doc("Đáp án và hướng dẫn chấm", subtitle)
    for row in assignments:
        q = row["assessment_question_versions"]
        number = row["display_number"]
        exam_doc.add_paragraph(f"Câu {number} ({row['assigned_score']} điểm). {q['prompt_text']}")
        for option in sorted(q.get("assessment_question_options") or [], key=lambda x: x["sequence_number"]):
            exam_doc.add_paragraph(f"{option['option_code']}. {option['option_text']}")
        for statement in sorted(q.get("assessment_question_statements") or [], key=lambda x: x["sequence_number"]):
            exam_doc.add_paragraph(f"{statement['statement_code']}. {statement['statement_text']}")
        answers = q.get("assessment_question_answers") or []
        if isinstance(answers, dict):
            answers = [answers]
        answer = answers[0] if answers else {}
        correct = [str(o["option_code"]) for o in q.get("assessment_question_options") or [] if o["is_correct"]]
        tf = [f"{s['statement_code']}: {'Đúng' if s['correct_value'] else 'Sai'}" for s in sorted(q.get("assessment_question_statements") or [], key=lambda x: x["sequence_number"])]
        solution = sorted(q.get("assessment_question_solutions") or [], key=lambda x: (not x["is_primary"], x["sequence_number"]))
        answer_doc.add_paragraph(f"Câu {number} ({row['assigned_score']} điểm): " + (str(answer.get("exact_answer_text") or answer.get("numeric_answer") or ", ".join(correct or tf) or "Xem lời giải")))
        if solution:
            answer_doc.add_paragraph(str(solution[0]["solution_text"]))
        elif answer.get("answer_explanation"):
            answer_doc.add_paragraph(str(answer["answer_explanation"]))

    matrix_doc = _doc("Ma trận đề kiểm tra", subtitle)
    _table(matrix_doc, ("Nội dung", "Mức độ", "Dạng", "Số câu", "Số ý", "Điểm"), [
        (c["topic_code"], c["cognitive_level_code"], c["question_type_code"], c["question_count"], c["response_count"], c["target_score"]) for c in cells
    ])
    spec_doc = _doc("Bản đặc tả đề kiểm tra", subtitle)
    spec_rows = []
    for row in assignments:
        q = row["assessment_question_versions"]
        c = cell_map[row["blueprint_cell_id"]]
        links = q.get("assessment_question_requirement_links") or []
        codes = [link["requirement_code"] for link in links if link["link_role"] == "PRIMARY"] or [link["requirement_code"] for link in links]
        descriptions = "; ".join(requirement_text.get(code, code) for code in codes)
        spec_rows.append((row["display_number"], c["topic_code"], descriptions or c["specification_note"], c["cognitive_level_code"], c["question_type_code"], row["assigned_score"]))
    _table(spec_doc, ("Câu", "Chủ đề", "Yêu cầu cần đạt", "Mức độ", "Dạng", "Điểm"), spec_rows)
    return {
        "de_kiem_tra_ban_nhap.docx": _bytes(exam_doc),
        "ma_tran_ban_nhap.docx": _bytes(matrix_doc),
        "ban_dac_ta_ban_nhap.docx": _bytes(spec_doc),
        "dap_an_huong_dan_cham_ban_nhap.docx": _bytes(answer_doc),
    }

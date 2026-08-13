"""Export an enriched lesson plan to a polished, editable Word document."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from models.lesson_plan_content import LessonPlanContent


BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
HEADER_FILL = "E8EEF5"
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGIN_DXA = {"top": 80, "bottom": 80, "start": 120, "end": 120}


def _set_run_font(run, size=11, bold=None, color=None):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def _shade_cell(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_cell_margins(cell):
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for side, value in CELL_MARGIN_DXA.items():
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_geometry(table, widths):
    table.autofit = False
    properties = table._tbl.tblPr
    width = properties.first_child_found_in("w:tblW")
    width.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    width.set(qn("w:type"), "dxa")
    indent = OxmlElement("w:tblInd")
    indent.set(qn("w:w"), str(TABLE_INDENT_DXA))
    indent.set(qn("w:type"), "dxa")
    properties.append(indent)

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for column_width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(column_width))
        grid.append(column)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index] / 1440)
            tc_width = cell._tc.get_or_add_tcPr().get_or_add_tcW()
            tc_width.set(qn("w:w"), str(widths[index]))
            tc_width.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell)


class LessonPlanDocxExporter:
    """Create a compact-reference-guide DOCX from LessonPlanContent."""

    def export(self, plan: LessonPlanContent, output: Path) -> Path:
        document = Document()
        self._configure_document(document)
        self._add_header(document, plan)
        self._add_objectives(document, plan)
        self._add_resources(document, plan)
        self._add_activities(document, plan)

        output.parent.mkdir(parents=True, exist_ok=True)
        document.save(output)
        return output

    @staticmethod
    def _configure_document(document):
        section = document.sections[0]
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.right_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)

        normal = document.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(11)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.25

        for name, size, color, before, after in (
            ("Heading 1", 16, BLUE, 18, 10),
            ("Heading 2", 13, BLUE, 14, 7),
            ("Heading 3", 12, DARK_BLUE, 10, 5),
        ):
            style = document.styles[name]
            style.font.name = "Calibri"
            style.font.size = Pt(size)
            style.font.color.rgb = color
            style.font.bold = True
            style.paragraph_format.space_before = Pt(before)
            style.paragraph_format.space_after = Pt(after)

    @staticmethod
    def _add_header(document, plan):
        title = document.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title.paragraph_format.space_after = Pt(8)
        run = title.add_run("KẾ HOẠCH BÀI DẠY")
        _set_run_font(run, size=20, bold=True, color=DARK_BLUE)

        subtitle = document.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle.paragraph_format.space_after = Pt(14)
        run = subtitle.add_run(plan.lesson_name)
        _set_run_font(run, size=14, bold=True)

        table = document.add_table(rows=2, cols=2)
        table.style = "Table Grid"
        values = (
            ("Môn học", plan.subject),
            ("Lớp", plan.grade),
            ("Số tiết", str(plan.total_periods)),
            ("Mẫu", str(plan.metadata.get("schema_id", ""))),
        )
        for cell, (label, value) in zip(
            [cell for row in table.rows for cell in row.cells], values
        ):
            paragraph = cell.paragraphs[0]
            label_run = paragraph.add_run(f"{label}: ")
            _set_run_font(label_run, bold=True)
            value_run = paragraph.add_run(value)
            _set_run_font(value_run)
        _set_table_geometry(table, [4680, 4680])

    @staticmethod
    def _add_bullet(document, text):
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing = 1.25
        _set_run_font(paragraph.add_run(text))

    def _add_objectives(self, document, plan):
        document.add_heading("I. MỤC TIÊU", level=1)
        for requirement in plan.objectives.knowledge:
            self._add_bullet(document, requirement)

    def _add_resources(self, document, plan):
        document.add_heading("II. THIẾT BỊ DẠY HỌC VÀ HỌC LIỆU", level=1)
        teacher = ", ".join(plan.resources.teacher) or "Không có dữ liệu"
        students = ", ".join(plan.resources.students) or "Không có dữ liệu"
        for label, value in (("Giáo viên", teacher), ("Học sinh", students)):
            paragraph = document.add_paragraph()
            _set_run_font(paragraph.add_run(f"{label}: "), bold=True)
            _set_run_font(paragraph.add_run(value))

    @staticmethod
    def _add_activities(document, plan):
        document.add_heading("III. TIẾN TRÌNH DẠY HỌC", level=1)
        for activity in plan.activities:
            activity_heading = document.add_heading(activity.title, level=2)
            activity_heading.paragraph_format.keep_with_next = True
            for label, value in (
                ("Mục tiêu", activity.objective),
                ("Nội dung", activity.content),
                ("Sản phẩm", activity.product),
            ):
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.keep_with_next = True
                _set_run_font(paragraph.add_run(f"{label}: "), bold=True)
                _set_run_font(paragraph.add_run(value))

            table = document.add_table(rows=1, cols=3)
            table.style = "Table Grid"
            headers = ("Bước", "Hoạt động của giáo viên và học sinh", "Sản phẩm")
            for cell, heading in zip(table.rows[0].cells, headers):
                _shade_cell(cell, HEADER_FILL)
                paragraph = cell.paragraphs[0]
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _set_run_font(paragraph.add_run(heading), bold=True)
            for step in activity.organization_steps:
                cells = table.add_row().cells
                values = (
                    step.title,
                    step.teacher_student_activity,
                    step.expected_product,
                )
                for cell, value in zip(cells, values):
                    _set_run_font(cell.paragraphs[0].add_run(value))
            _set_table_geometry(table, [1800, 4860, 2700])

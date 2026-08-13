from pathlib import Path

from docx import Document

from exporters import LessonPlanDocxExporter
from intelligence.lesson_plan_builder import LessonPlanBuilder
from intelligence.lesson_plan_content_enricher import LessonPlanContentEnricher
from models.lesson_model import LessonModel
from models.math_lesson_plan_schema import create_math_lesson_plan_schema


def make_plan():
    lesson = LessonModel(
        subject="Toán",
        grade="6",
        lesson_name="Phân số với tử và mẫu là số nguyên",
        period_count=2,
        learning_requirements=["Nhận biết được phân số."],
        registered_equipment=["Máy chiếu"],
        learning_resources=["Vở ghi"],
    )
    plan = LessonPlanBuilder().build(
        lesson=lesson,
        schema=create_math_lesson_plan_schema(),
    )
    return LessonPlanContentEnricher().enrich(plan)


def test_docx_exporter_creates_editable_lesson_plan(tmp_path):
    output = tmp_path / "lesson-plan.docx"
    LessonPlanDocxExporter().export(make_plan(), output)

    assert output.exists()
    document = Document(output)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "KẾ HOẠCH BÀI DẠY" in text
    assert "Phân số với tử và mẫu là số nguyên" in text
    assert "I. MỤC TIÊU" in text
    assert "III. TIẾN TRÌNH DẠY HỌC" in text
    assert len(document.tables) == 5


def test_docx_uses_expected_page_geometry(tmp_path):
    output = tmp_path / "lesson-plan.docx"
    LessonPlanDocxExporter().export(make_plan(), output)
    section = Document(output).sections[0]

    assert round(section.page_width.inches, 1) == 8.5
    assert round(section.page_height.inches, 1) == 11.0
    assert round(section.left_margin.inches, 1) == 1.0
    assert round(section.right_margin.inches, 1) == 1.0

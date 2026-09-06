from __future__ import annotations
from datetime import date
from io import BytesIO
from docx import Document
from document_standardization.lesson_plan_multi_period_scope import (
    PeriodScopeStatus, apply_scoped_english_period_dates, resolve_multi_period_scope,
)

def _docx_bytes():
    doc=Document()
    for period in (1,2):
        table=doc.add_table(rows=1,cols=1)
        table.cell(0,0).text=(
            "Date of planning: 04/9/2023\n"
            "Date of teaching: /9/2023\n"
            f"Peroid {period} : UNIT 1: MY NEW SCHOOL"
        )
    b=BytesIO();doc.save(b);return b.getvalue()

def _context():
    return {"curriculum_periods":(1,2),"occurrences":(
        {"curriculum_period":1,"teaching_date":date(2026,9,7),"class_id":"6A1","class_display":"6A1","timetable_period":1},
        {"curriculum_period":2,"teaching_date":date(2026,9,8),"class_id":"6A1","class_display":"6A1","timetable_period":2},
    )}

def _text(content):
    doc=Document(BytesIO(content))
    return "\n".join(p.text for t in doc.tables for r in t.rows for c in r.cells for p in c.paragraphs)

def test_legacy_peroid_and_incomplete_teaching_date_are_fully_replaced():
    source=_docx_bytes()
    scope=resolve_multi_period_scope(group_context=_context(),document_periods=(1,2))
    assert scope.status == PeriodScopeStatus.SUCCESS
    result=apply_scoped_english_period_dates(
        source_content=source,output_content=source,scope=scope,drafting_date=date(2026,9,4)
    )
    assert result.applied_periods == (1,2)
    text=_text(result.content)
    assert "04/09/2026" in text
    assert "07/09/2026" in text
    assert "08/09/2026" in text
    assert "/9/2023" not in text
    assert "07/09/2026/9/2023" not in text
    assert "08/09/2026/9/2023" not in text

def test_patch_preserves_table_count():
    source=_docx_bytes()
    before=Document(BytesIO(source))
    scope=resolve_multi_period_scope(group_context=_context(),document_periods=(1,2))
    result=apply_scoped_english_period_dates(
        source_content=source,output_content=source,scope=scope,drafting_date=date(2026,9,4)
    )
    after=Document(BytesIO(result.content))
    assert len(after.tables)==len(before.tables)==2